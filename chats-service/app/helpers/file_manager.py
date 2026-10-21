"""Document text extraction and page-aware chunking."""

import logging
import bisect
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path

import aiofiles
import docx
import openpyxl
import pymupdf

from app.core.config import settings
from app.helpers.storage import download_file

logger = logging.getLogger(__name__)


@dataclass
class Page:
    """One unit of extracted text.

    number is 1-based for paged formats and 0 for everything else, so a
    citation never claims "page 0".
    """

    number: int
    text: str
    section: str | None = None


@dataclass
class Chunk:
    index: int
    text: str  # stored as document_chunks.chunk_text
    embed_text: str  # what is actually embedded; may carry a context header
    page: int | None
    section: str | None


def _outline_by_page(doc) -> dict[int, str]:
    """Section title per page, taken from the PDF outline.

    The usual alternative is guessing headings from font sizes, which is
    unreliable. A bookmark tree, when the document has one, is authoritative.
    """
    sections: dict[int, str] = {}
    try:
        for entry in doc.get_toc() or []:
            _level, title, page = entry[0], entry[1], entry[2]
            if page and page >= 1 and title.strip():
                sections.setdefault(page, title.strip())
    except Exception:
        # An absent or malformed outline is not a reason to fail extraction.
        pass
    return sections


def _extract_pdf(path: str) -> list[Page]:
    pages: list[Page] = []
    with pymupdf.open(path) as doc:
        sections = _outline_by_page(doc)
        current = None
        for number, page in enumerate(doc, start=1):
            # A heading applies until the next one, so carry it forward.
            current = sections.get(number, current)
            pages.append(Page(number, page.get_text("text"), current))
    return pages


def _extract_docx(path: str) -> list[Page]:
    """One Page per heading-delimited section.

    A .docx has no page numbers until it is laid out, so chunks from one carry
    a section but no page.
    """
    document = docx.Document(path)
    pages: list[Page] = []
    heading: str | None = None
    buffer: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style = getattr(paragraph.style, "name", "") or ""
        if style.startswith("Heading"):
            if buffer:
                pages.append(Page(0, "\n".join(buffer), heading))
                buffer = []
            heading = text
        else:
            buffer.append(text)

    if buffer:
        pages.append(Page(0, "\n".join(buffer), heading))
    return pages


def _extract_xlsx(path: str) -> list[Page]:
    """One Page per worksheet, the sheet name as its section."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        pages: list[Page] = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                pages.append(Page(0, "\n".join(rows), sheet.title))
        return pages
    finally:
        workbook.close()


def _extract_txt(path: str) -> list[Page]:
    with open(path, "r", encoding="utf-8", errors="ignore") as handle:
        return [Page(0, handle.read())]


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".xlsm": _extract_xlsx,
    ".txt": _extract_txt,
    ".md": _extract_txt,
    ".csv": _extract_txt,
}

# Legacy binary formats nothing here can read. Named explicitly so the error
# says what to do instead of "unsupported".
_LEGACY = {".doc": ".docx", ".xls": ".xlsx"}


_DOT_LEADER = re.compile(r"\.{3,}")


def _is_navigational(text: str) -> bool:
    """True for table-of-contents and index pages.

    They are dot leaders and page numbers: they name everything in the
    document and explain none of it, so they match almost any keyword query
    while carrying no information. Keyed on dot-leader density specifically
    rather than a general punctuation ratio, which would also catch code.
    """
    if not text.strip():
        return True

    leader_chars = sum(len(m.group()) for m in _DOT_LEADER.finditer(text))
    return leader_chars / len(text) > 0.15


def _context_header(document_title: str, section: str | None, page: int | None) -> str:
    label = " > ".join(p for p in (document_title, section) if p)
    if page:
        label = f"{label} (page {page})"
    return f"{label}\n\n" if label else ""


class FileManager:
    """Downloads a stored document and turns it into embeddable chunks."""

    @staticmethod
    async def download_to_tempfile(key: str) -> str:
        content = await download_file(key)
        suffix = Path(key).suffix or ".bin"

        handle = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
        handle.close()
        async with aiofiles.open(handle.name, "wb") as f:
            await f.write(content)
        return handle.name

    @staticmethod
    def extract_pages(file_path: str) -> list[Page]:
        ext = Path(file_path).suffix.lower()

        if ext in _LEGACY:
            raise ValueError(
                f"{ext} is a legacy binary format and cannot be read; "
                f"convert it to {_LEGACY[ext]}"
            )

        extractor = _EXTRACTORS.get(ext)
        if extractor is None:
            raise ValueError(f"Unsupported file type: {ext}")

        try:
            return extractor(file_path)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Text extraction failed for {ext}: {exc}") from exc

    @staticmethod
    def chunk_pages(
        pages: list[Page],
        document_title: str,
        *,
        chunk_size: int | None = None,
        overlap: int | None = None,
    ) -> list[Chunk]:
        """Split the document into overlapping chunks that know where they came from.

        Pages are concatenated and chunked as one stream rather than page by
        page: chunking per page would fragment at every page break. Each page's
        start offset is recorded, so a chunk is attributed to the page its text
        begins on.
        """
        chunk_size = chunk_size or settings.CHUNK_SIZE
        overlap = overlap or settings.CHUNK_OVERLAP
        step = max(1, chunk_size - overlap)

        pieces: list[str] = []
        marks: list[tuple[int, int, str | None]] = []
        cursor = 0

        for page in pages:
            text = " ".join(page.text.split())
            if not text:
                continue
            marks.append((cursor, page.number, page.section))
            pieces.append(text)
            cursor += len(text) + 1  # +1 for the joining space

        full = " ".join(pieces)
        if not full:
            return []

        starts = [m[0] for m in marks]
        chunks: list[Chunk] = []
        skipped = 0
        position = 0

        while position < len(full):
            raw = full[position : position + chunk_size].strip()

            if raw and settings.SKIP_NAVIGATIONAL_CHUNKS and _is_navigational(raw):
                skipped += 1
                raw = ""

            if raw:
                at = bisect.bisect_right(starts, position) - 1
                _, number, section = marks[at] if at >= 0 else (0, 0, None)
                header = (
                    _context_header(document_title, section, number or None)
                    if settings.CONTEXTUAL_CHUNK_HEADER
                    else ""
                )
                chunks.append(
                    Chunk(
                        index=len(chunks),
                        text=raw,
                        embed_text=f"{header}{raw}",
                        page=number or None,
                        section=section,
                    )
                )
            position += step

        if skipped:
            logger.info(f"[CHUNK] {document_title}: skipped {skipped} navigational chunk(s)")

        return chunks

    @staticmethod
    async def get_chunks(source: dict, document_title: str) -> list[Chunk]:
        """Download, extract and chunk one stored document."""
        if not isinstance(source, dict) or "s3_key" not in source:
            raise ValueError("Invalid source: expected a dict with 's3_key'")

        tmp_path = await FileManager.download_to_tempfile(source["s3_key"])
        try:
            pages = FileManager.extract_pages(tmp_path)
        finally:
            # The previous implementation never removed these.
            try:
                os.unlink(tmp_path)
            except OSError:
                pass

        return FileManager.chunk_pages(pages, document_title)
